import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pathlib import Path
from datetime import datetime
import json
import os
import sys
import threading
import queue
import traceback
import subprocess

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Border, Side

try:
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
    MATPLOTLIB_AVAILABLE = True
except Exception:
    MATPLOTLIB_AVAILABLE = False

from burnin_runtime import make_burnin_config, run_burnin, BurnInError


# ---------------- Persistent settings (default folder) ----------------

CONFIG_FILE = Path(__file__).with_name("config.json")
PERCENT_THRESHOLD = 10.0
GUI_AVERAGE_SAMPLES_PER_CYCLE = 10


def load_config() -> dict:
    if CONFIG_FILE.exists():
        try:
            return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_config(cfg: dict) -> None:
    CONFIG_FILE.write_text(json.dumps(cfg, indent=2), encoding="utf-8")


# ---------------- Excel formatting ----------------

DARK_GREEN_FILL = PatternFill(fill_type="solid", start_color="008000", end_color="008000")
LIGHT_GREEN_FILL = PatternFill(fill_type="solid", start_color="92D050", end_color="92D050")
MID_GREEN_FILL = PatternFill(fill_type="solid", start_color="C6E0B4", end_color="C6E0B4")
ORANGE_FILL = PatternFill(fill_type="solid", start_color="F4B183", end_color="F4B183")
RED_FILL = PatternFill(fill_type="solid", start_color="FF6666", end_color="FF6666")

THIN_BLACK_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


def get_resistance_fill(value: float):
    if value <= 0.1:
        return DARK_GREEN_FILL
    elif value <= 1.0:
        return LIGHT_GREEN_FILL
    elif value <= 3000000.0:
        return ORANGE_FILL
    else:
        return RED_FILL


def make_workbook_name(serial: str) -> str:
    serial_clean = serial.strip()
    return f"{serial_clean}.xlsx"


def safe_sheet_title(title: str) -> str:
    """Excel sheet names max out at 31 chars and disallow certain characters."""
    bad_chars = '[]:*?/\\'
    cleaned = "".join("_" if ch in bad_chars else ch for ch in title)
    return cleaned[:31]


def open_file(path: Path) -> None:
    try:
        path = str(path)
        if sys.platform.startswith("win"):
            os.startfile(path)
        elif sys.platform.startswith("darwin"):
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])
    except Exception:
        pass


def export_run_to_excel(
    outdir: str,
    serial: str,
    phase: str,
    tester_name: str,
    notes: str,
    data_2d: list[list],
) -> Path:
    out_path = Path(outdir).expanduser().resolve()
    out_path.mkdir(parents=True, exist_ok=True)

    wb_path = out_path / make_workbook_name(serial)

    if wb_path.exists():
        wb = load_workbook(wb_path)
    else:
        wb = Workbook()
        if wb.sheetnames == ["Sheet"]:
            wb.remove(wb["Sheet"])

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    phase_clean = phase.strip().replace(" ", "_")
    ws = wb.create_sheet(title=safe_sheet_title(f"{phase_clean}_{ts}"))

    meta_rows = [
        ("Serial Number", serial.strip().upper()),
        ("Testing Phase", phase),
        ("Tester Name", tester_name),
        ("Notes", notes if notes.strip() else "(none)"),
        ("Run Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]

    for i, (k, v) in enumerate(meta_rows, start=1):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)

    data_start_row = len(meta_rows) + 10 + 1

    for r_i, row in enumerate(data_2d):
        for c_i, val in enumerate(row):
            cell = ws.cell(row=data_start_row + r_i, column=1 + c_i, value=val)

            if val is not None:
                cell.border = THIN_BLACK_BORDER

            if isinstance(val, (int, float)) and not (r_i == 0 or c_i == 0):
                cell.number_format = '0.00E+00'
                cell.fill = get_resistance_fill(val)

    max_cols = 2
    for row in data_2d:
        if isinstance(row, list):
            max_cols = max(max_cols, len(row))

    for col in range(1, max_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    wb.save(wb_path)
    return wb_path


def export_burnin_to_excel(
    *,
    outdir: str,
    serial: str,
    phase: str,
    tester_name: str,
    notes: str,
    burnin_config,
    rows: list[dict],
    final_status: str,
    failure_reason: str = "",
) -> Path:
    out_path = Path(outdir).expanduser().resolve()
    out_path.mkdir(parents=True, exist_ok=True)
    wb_path = out_path / make_workbook_name(serial)

    if wb_path.exists():
        wb = load_workbook(wb_path)
    else:
        wb = Workbook()
        if wb.sheetnames == ["Sheet"]:
            wb.remove(wb["Sheet"])

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    phase_clean = phase.strip().replace(" ", "_")
    # User requested the same style as Resistance_Map_{testing phase metadata}.
    ws = wb.create_sheet(title=safe_sheet_title(f"Resistance_Map_{phase_clean}_{ts}"))

    baseline = rows[0].get("baseline_ohms", "") if rows else ""

    meta_rows = [
        ("Serial Number", serial.strip().upper()),
        ("Testing Phase", phase),
        ("Tester Name", tester_name),
        ("Notes", notes if notes.strip() else "(none)"),
        ("Pad 1", burnin_config.pad1),
        ("Pad 2", burnin_config.pad2),
        ("Current (A, metadata / external supply)", burnin_config.current_amps),
        ("On Time (s)", burnin_config.on_time),
        ("Off Time (s)", burnin_config.off_time),
        ("Total Cycles", burnin_config.pulses),
        ("Resistance Threshold", f"±{burnin_config.percent_threshold:.2f}%"),
        ("Baseline Resistance (Ω)", baseline),
        ("Final Status", final_status),
        ("Failure Reason", failure_reason if failure_reason else "(none)"),
        ("Run Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]

    for i, (k, v) in enumerate(meta_rows, start=1):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)

    data_start_row = len(meta_rows) + 10 + 1
    headers = [
        "Cycle #",
        "Sample #",
        "Time Since Start (s)",
        "Cycle Elapsed Time (s)",
        "Resistance (Ω)",
        "Baseline Resistance (Ω)",
        "Percent Change (%)",
        "Status",
    ]

    for c_i, header in enumerate(headers, start=1):
        cell = ws.cell(row=data_start_row, column=c_i, value=header)
        cell.border = THIN_BLACK_BORDER

    for r_i, row in enumerate(rows, start=data_start_row + 1):
        values = [
            row.get("cycle"),
            row.get("sample_index"),
            row.get("elapsed_s"),
            row.get("cycle_elapsed_s"),
            row.get("resistance_ohms"),
            row.get("baseline_ohms"),
            row.get("percent_change"),
            row.get("status"),
        ]
        for c_i, value in enumerate(values, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=value)
            cell.border = THIN_BLACK_BORDER
            if isinstance(value, (int, float)):
                if c_i in (5, 6):
                    cell.number_format = "0.000000"
                    cell.fill = get_resistance_fill(float(value))
                elif c_i == 7:
                    cell.number_format = "0.00"
                else:
                    cell.number_format = "0.000"

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    wb.save(wb_path)
    return wb_path


# ---------------- Shared metadata panel ----------------

class MetadataPanel(ttk.LabelFrame):
    def __init__(self, master, serial_var, phase_var, tester_var, output_dir_var, app, title="Metadata"):
        super().__init__(master, text=title, padding=10)

        self.serial_var = serial_var
        self.phase_var = phase_var
        self.tester_var = tester_var
        self.output_dir_var = output_dir_var
        self.app = app

        self.columnconfigure(1, weight=1)

        ttk.Label(self, text="Serial Number (SNO###):").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=4)
        self.serial_combo = ttk.Combobox(self, textvariable=self.serial_var)
        self.serial_combo.grid(row=0, column=1, sticky="ew", pady=4)
        self.serial_combo["values"] = ()

        ttk.Label(self, text="Testing Phase:").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=4)
        self.phase_entry = ttk.Entry(self, textvariable=self.phase_var)
        self.phase_entry.grid(row=1, column=1, sticky="ew", pady=4)

        ttk.Label(self, text="Tester Name:").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=4)
        self.tester_entry = ttk.Entry(self, textvariable=self.tester_var)
        self.tester_entry.grid(row=2, column=1, sticky="ew", pady=4)

        ttk.Label(self, text="Other Notes:").grid(row=3, column=0, sticky="nw", padx=(0, 8), pady=4)
        self.notes_text = tk.Text(self, height=5, wrap="word")
        self.notes_text.grid(row=3, column=1, sticky="ew", pady=4)

        ttk.Label(self, text="Default Save Folder:").grid(row=4, column=0, sticky="w", padx=(0, 8), pady=4)
        self.output_entry = ttk.Entry(self, textvariable=self.output_dir_var)
        self.output_entry.grid(row=4, column=1, sticky="ew", pady=4)

        ttk.Button(self, text="Change…", command=self.app.change_default_folder).grid(
            row=4, column=2, sticky="e", padx=(8, 0), pady=4
        )

        ttk.Button(self, text="Refresh SN List", command=self.app.refresh_serial_list).grid(
            row=4, column=3, sticky="e", padx=(8, 0), pady=4
        )

    def get_notes(self) -> str:
        return self.notes_text.get("1.0", "end").strip()

    def set_serial_values(self, serials: list[str]) -> None:
        self.serial_combo["values"] = serials


# ---------------- Main Tkinter GUI ----------------

class App(ttk.Frame):
    def __init__(self, master: tk.Tk):
        super().__init__(master, padding=12)
        self.master = master

        cfg = load_config()
        saved_dir = cfg.get("default_output_dir", "")
        default_dir = saved_dir if saved_dir and Path(saved_dir).exists() else ""

        self.serial_var = tk.StringVar()
        self.phase_var = tk.StringVar()
        self.tester_var = tk.StringVar()
        self.output_dir_var = tk.StringVar(value=default_dir)

        self.resistance_status_var = tk.StringVar(value="Ready.")
        self.burnin_status_var = tk.StringVar(value="Ready.")

        self.burnin_on_time_var = tk.StringVar()
        self.burnin_off_time_var = tk.StringVar()
        self.burnin_pad_1_var = tk.StringVar()
        self.burnin_pad_2_var = tk.StringVar()
        self.burnin_current_var = tk.StringVar()
        self.burnin_total_cycles_var = tk.StringVar()

        self.metadata_panels: list[MetadataPanel] = []

        self.burnin_queue = queue.Queue()
        self.burnin_thread = None
        self.burnin_stop_event = threading.Event()
        self.burnin_running = False
        # Raw rows are every measurement streamed by burnin_runtime.
        # Display/export rows are GUI-reduced to one averaged point per cycle.
        self.burnin_raw_rows: list[dict] = []
        self.burnin_rows: list[dict] = []
        self.burnin_cycle_buffers: dict[int, list[dict]] = {}
        self.burnin_completed_cycles: set[int] = set()
        self.burnin_x: list[float] = []
        self.burnin_y: list[float] = []
        self.burnin_config = None
        self.burnin_final_status = ""
        self.burnin_failure_reason = ""

        self.grid(row=0, column=0, sticky="nsew")
        master.columnconfigure(0, weight=1)
        master.rowconfigure(0, weight=1)
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)

        self._build_tabs()
        self._bind_shortcuts()

        self.after(150, self.ensure_output_folder)
        self.after(200, self.refresh_serial_list)
        self.after(100, self.process_burnin_queue)

    # ---------------- UI BUILDERS ----------------
    def _build_tabs(self) -> None:
        self.notebook = ttk.Notebook(self)
        self.notebook.grid(row=0, column=0, sticky="nsew")

        self.resistance_tab = ttk.Frame(self.notebook, padding=10)
        self.burnin_tab = ttk.Frame(self.notebook, padding=10)

        self.notebook.add(self.resistance_tab, text="Resistance Mapping")
        self.notebook.add(self.burnin_tab, text="Burn In")

        self._build_resistance_tab()
        self._build_burnin_tab()

    def _build_resistance_tab(self) -> None:
        tab = self.resistance_tab
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(2, weight=1)

        self.resistance_metadata = MetadataPanel(tab, self.serial_var, self.phase_var, self.tester_var, self.output_dir_var, self)
        self.resistance_metadata.grid(row=0, column=0, sticky="ew")
        self.metadata_panels.append(self.resistance_metadata)
        self.resistance_metadata.serial_combo.focus_set()

        self._build_resistance_controls(tab)
        self._build_resistance_log(tab)

    def _build_resistance_controls(self, parent) -> None:
        lf = ttk.Frame(parent)
        lf.grid(row=1, column=0, sticky="ew", pady=(10, 0))
        lf.columnconfigure(0, weight=1)

        ttk.Label(lf, textvariable=self.resistance_status_var).grid(row=0, column=0, sticky="w")
        btns = ttk.Frame(lf)
        btns.grid(row=0, column=1, sticky="e")

        self.resistance_start_btn = ttk.Button(btns, text="Start", command=self.on_resistance_start)
        self.resistance_start_btn.grid(row=0, column=0, padx=(0, 8))
        ttk.Button(btns, text="Stop", command=self.on_resistance_stop).grid(row=0, column=1, padx=(0, 8))
        ttk.Button(btns, text="Clear Log", command=self.clear_resistance_log).grid(row=0, column=2)

    def _build_resistance_log(self, parent) -> None:
        lf = ttk.LabelFrame(parent, text="Log", padding=10)
        lf.grid(row=2, column=0, sticky="nsew", pady=(10, 0))
        lf.rowconfigure(0, weight=1)
        lf.columnconfigure(0, weight=1)

        self.resistance_log_text = tk.Text(lf, height=10, wrap="word", state="disabled")
        self.resistance_log_text.grid(row=0, column=0, sticky="nsew")

        scrollbar = ttk.Scrollbar(lf, orient="vertical", command=self.resistance_log_text.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.resistance_log_text.configure(yscrollcommand=scrollbar.set)

    def _build_burnin_tab(self) -> None:
        tab = self.burnin_tab
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(2, weight=1)
        tab.rowconfigure(4, weight=1)

        self.burnin_metadata = MetadataPanel(tab, self.serial_var, self.phase_var, self.tester_var, self.output_dir_var, self)
        self.burnin_metadata.grid(row=0, column=0, sticky="ew")
        self.metadata_panels.append(self.burnin_metadata)

        self._build_burnin_inputs(tab)
        self._build_burnin_graph(tab)
        self._build_burnin_controls(tab)
        self._build_burnin_log(tab)

    def _build_burnin_inputs(self, parent) -> None:
        lf = ttk.LabelFrame(parent, text="Burn-In Settings (Optional except pads)", padding=10)
        lf.grid(row=1, column=0, sticky="ew", pady=(10, 0))

        for col in range(6):
            lf.columnconfigure(col, weight=1)

        ttk.Label(lf, text="On Time (s):").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(lf, textvariable=self.burnin_on_time_var, width=14).grid(row=0, column=1, sticky="ew", pady=4)

        ttk.Label(lf, text="Off Time (s):").grid(row=0, column=2, sticky="w", padx=(12, 8), pady=4)
        ttk.Entry(lf, textvariable=self.burnin_off_time_var, width=14).grid(row=0, column=3, sticky="ew", pady=4)

        ttk.Label(lf, text="Current (A):").grid(row=0, column=4, sticky="w", padx=(12, 8), pady=4)
        ttk.Entry(lf, textvariable=self.burnin_current_var, width=14).grid(row=0, column=5, sticky="ew", pady=4)

        ttk.Label(lf, text="Pad 1:").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(lf, textvariable=self.burnin_pad_1_var, width=14).grid(row=1, column=1, sticky="ew", pady=4)

        ttk.Label(lf, text="Pad 2:").grid(row=1, column=2, sticky="w", padx=(12, 8), pady=4)
        ttk.Entry(lf, textvariable=self.burnin_pad_2_var, width=14).grid(row=1, column=3, sticky="ew", pady=4)

        ttk.Label(lf, text="Total Cycles (#):").grid(row=1, column=4, sticky="w", padx=(12, 8), pady=4)
        ttk.Entry(lf, textvariable=self.burnin_total_cycles_var, width=14).grid(row=1, column=5, sticky="ew", pady=4)

        ttk.Label(
            lf,
            text="Blank on/off/current/cycles use burninConfig.json defaults. Threshold is hardcoded at ±10%.",
            foreground="gray",
        ).grid(row=2, column=0, columnspan=6, sticky="w", pady=(6, 0))

    def _build_burnin_graph(self, parent) -> None:
        lf = ttk.LabelFrame(parent, text="Live Burn-In Graph", padding=10)
        lf.grid(row=2, column=0, sticky="nsew", pady=(10, 0))
        lf.rowconfigure(0, weight=1)
        lf.columnconfigure(0, weight=1)

        if MATPLOTLIB_AVAILABLE:
            self.burnin_fig = Figure(figsize=(7, 3.2), dpi=100)
            self.burnin_ax = self.burnin_fig.add_subplot(111)
            self.burnin_line, = self.burnin_ax.plot([], [], marker="o")
            self.burnin_ax.set_xlabel("Cycle #")
            self.burnin_ax.set_ylabel("Resistance (Ω)")
            self.burnin_ax.set_title("Burn-In Resistance")
            self.burnin_ax.grid(True)
            self.burnin_canvas = FigureCanvasTkAgg(self.burnin_fig, master=lf)
            self.burnin_canvas.get_tk_widget().grid(row=0, column=0, sticky="nsew")
        else:
            self.burnin_graph_canvas = tk.Canvas(lf, height=260, background="white", highlightthickness=1, highlightbackground="#cccccc")
            self.burnin_graph_canvas.grid(row=0, column=0, sticky="nsew")
            self.burnin_graph_canvas.bind("<Configure>", self._draw_burnin_graph_placeholder)

    def _build_burnin_controls(self, parent) -> None:
        lf = ttk.Frame(parent)
        lf.grid(row=3, column=0, sticky="ew", pady=(10, 0))
        lf.columnconfigure(0, weight=1)

        ttk.Label(lf, textvariable=self.burnin_status_var).grid(row=0, column=0, sticky="w")
        btns = ttk.Frame(lf)
        btns.grid(row=0, column=1, sticky="e")

        self.burnin_start_btn = ttk.Button(btns, text="Start Burn-In", command=self.on_burnin_start)
        self.burnin_start_btn.grid(row=0, column=0, padx=(0, 8))
        self.burnin_stop_btn = ttk.Button(btns, text="Stop", command=self.on_burnin_stop, state="disabled")
        self.burnin_stop_btn.grid(row=0, column=1, padx=(0, 8))
        ttk.Button(btns, text="Clear Log", command=self.clear_burnin_log).grid(row=0, column=2)

    def _build_burnin_log(self, parent) -> None:
        lf = ttk.LabelFrame(parent, text="Burn-In Log", padding=10)
        lf.grid(row=4, column=0, sticky="nsew", pady=(10, 0))
        lf.rowconfigure(0, weight=1)
        lf.columnconfigure(0, weight=1)

        self.burnin_log_text = tk.Text(lf, height=7, wrap="word", state="disabled")
        self.burnin_log_text.grid(row=0, column=0, sticky="nsew")

        scrollbar = ttk.Scrollbar(lf, orient="vertical", command=self.burnin_log_text.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.burnin_log_text.configure(yscrollcommand=scrollbar.set)

    def _bind_shortcuts(self) -> None:
        self.master.bind("<Control-Return>", lambda e: self.start_active_tab())
        self.master.bind("<Escape>", lambda e: self.stop_active_tab())

    # ---------------- COMMON HELPERS ----------------
    def start_active_tab(self) -> None:
        selected = self.notebook.select()
        if selected == str(self.resistance_tab):
            self.on_resistance_start()
        elif selected == str(self.burnin_tab):
            self.on_burnin_start()

    def stop_active_tab(self) -> None:
        selected = self.notebook.select()
        if selected == str(self.resistance_tab):
            self.on_resistance_stop()
        elif selected == str(self.burnin_tab):
            self.on_burnin_stop()

    def log_to_widget(self, widget: tk.Text, msg: str) -> None:
        timestamp = datetime.now().strftime("%H:%M:%S")
        line = f"[{timestamp}] {msg}\n"
        widget.configure(state="normal")
        widget.insert("end", line)
        widget.see("end")
        widget.configure(state="disabled")

    def resistance_log(self, msg: str) -> None:
        self.log_to_widget(self.resistance_log_text, msg)

    def burnin_log(self, msg: str) -> None:
        self.log_to_widget(self.burnin_log_text, msg)

    def clear_text_widget(self, widget: tk.Text) -> None:
        widget.configure(state="normal")
        widget.delete("1.0", "end")
        widget.configure(state="disabled")

    def clear_resistance_log(self) -> None:
        self.clear_text_widget(self.resistance_log_text)
        self.resistance_status_var.set("Ready.")

    def clear_burnin_log(self) -> None:
        self.clear_text_widget(self.burnin_log_text)
        self.burnin_status_var.set("Ready.")

    def ensure_output_folder(self) -> None:
        if not self.output_dir_var.get().strip():
            messagebox.showinfo(
                "Select Default Folder",
                "Please choose the default folder where all Excel files will be saved.\n"
                "This will be remembered next time you open the program."
            )
            self.change_default_folder(force=True)

    def change_default_folder(self, force: bool = False) -> None:
        current = self.output_dir_var.get().strip()
        initial_dir = current if current and Path(current).exists() else str(Path.home())

        folder = filedialog.askdirectory(title="Choose default output folder", initialdir=initial_dir, mustexist=True)

        if folder:
            self.output_dir_var.set(folder)
            cfg = load_config()
            cfg["default_output_dir"] = folder
            save_config(cfg)
            self.resistance_log(f"Default output folder set to: {folder}")
            self.burnin_log(f"Default output folder set to: {folder}")
            self.resistance_status_var.set("Default folder saved.")
            self.burnin_status_var.set("Default folder saved.")
            self.refresh_serial_list()

        elif force:
            messagebox.showwarning("Folder Required", "You must select a default output folder to continue.")
            self.after(150, self.ensure_output_folder)

    def get_existing_serials(self) -> list[str]:
        outdir = self.output_dir_var.get().strip()
        if not outdir:
            return []

        path = Path(outdir)
        if not path.exists() or not path.is_dir():
            return []

        serials = []
        for file in path.glob("*.xlsx"):
            stem = file.stem.strip().upper()
            if stem:
                serials.append(stem)

        return sorted(set(serials))

    def refresh_serial_list(self) -> None:
        serials = self.get_existing_serials()
        for panel in self.metadata_panels:
            panel.set_serial_values(serials)

        current = self.serial_var.get().strip().upper()
        if current:
            self.serial_var.set(current)

        self.resistance_log(f"Loaded {len(serials)} existing serial number(s) from folder.")
        self.burnin_log(f"Loaded {len(serials)} existing serial number(s) from folder.")

    def validate_common_inputs(self) -> tuple[bool, str]:
        serial = self.serial_var.get().strip().upper()
        phase = self.phase_var.get().strip()
        tester = self.tester_var.get().strip()
        outdir = self.output_dir_var.get().strip()

        if not serial:
            return False, "Serial Number is required."
        if not phase:
            return False, "Testing Phase is required."
        if not tester:
            return False, "Tester Name is required."
        if not outdir:
            return False, "Default output folder is not set."
        if not Path(outdir).exists():
            return False, "Default output folder does not exist."
        return True, ""

    def get_burnin_inputs(self) -> dict:
        return {
            "on_time": self.burnin_on_time_var.get().strip() or None,
            "off_time": self.burnin_off_time_var.get().strip() or None,
            "pad1": self.burnin_pad_1_var.get().strip(),
            "pad2": self.burnin_pad_2_var.get().strip(),
            "current_amps": self.burnin_current_var.get().strip() or None,
            "pulses": self.burnin_total_cycles_var.get().strip() or None,
        }

    def _draw_burnin_graph_placeholder(self, event=None) -> None:
        canvas = self.burnin_graph_canvas
        canvas.delete("all")
        width = canvas.winfo_width()
        height = canvas.winfo_height()
        pad = 40
        canvas.create_line(pad, height - pad, width - pad, height - pad, width=2)
        canvas.create_line(pad, height - pad, pad, pad, width=2)
        canvas.create_text(width / 2, height / 2, text="Matplotlib not available", fill="gray")
        canvas.create_text(width / 2, height - 12, text="Cycle #", fill="gray")
        canvas.create_text(14, height / 2, text="Resistance (Ω)", fill="gray", angle=90)

    # ---------------- RESISTANCE CALLBACKS ----------------
    def on_resistance_start(self) -> None:
        ok, err = self.validate_common_inputs()
        if not ok:
            messagebox.showerror("Missing / Invalid Info", err)
            self.resistance_status_var.set("Fix inputs and try again.")
            return

        serial = self.serial_var.get().strip().upper()
        phase = self.phase_var.get().strip()
        tester = self.tester_var.get().strip()
        notes = self.resistance_metadata.get_notes()
        outdir = self.output_dir_var.get().strip()

        self.resistance_status_var.set("Running...")
        self.resistance_log("Start pressed.")
        self.resistance_log(f"Serial: {serial} | Phase: {phase} | Tester: {tester}")

        try:
            from Resistance_map import run_resistance_map
            results_2d = run_resistance_map()
        except Exception as e:
            self.resistance_log(f"ERROR running resistance map: {e}")
            self.resistance_status_var.set("Resistance map failed.")
            messagebox.showerror(
                "Resistance Map Failed",
                f"Could not run resistance mapping.\n\n"
                f"Make sure the resistance-map PCB is connected .\n\nError:\n{e}"
            )
            return
        
        try:
            wb_path = export_run_to_excel(outdir, serial, phase, tester, notes, results_2d)
            self.resistance_log(f"Exported run to: {wb_path}")
            self.resistance_status_var.set("Finished.")
            open_file(wb_path)
            self.refresh_serial_list()
        except Exception as e:
            self.resistance_log(f"ERROR exporting: {e}")
            self.resistance_status_var.set("Export failed.")
            messagebox.showerror("Export Failed", str(e))

    def on_resistance_stop(self) -> None:
        self.resistance_status_var.set("Stopped (placeholder).")
        self.resistance_log("Stop pressed (placeholder).")

    # ---------------- BURN-IN CALLBACKS ----------------
    def on_burnin_start(self) -> None:
        if self.burnin_running:
            return

        ok, err = self.validate_common_inputs()
        if not ok:
            messagebox.showerror("Missing / Invalid Info", err)
            self.burnin_status_var.set("Fix inputs and try again.")
            return

        try:
            self.burnin_config = make_burnin_config(**self.get_burnin_inputs())
        except Exception as exc:
            messagebox.showerror("Invalid Burn-In Settings", str(exc))
            self.burnin_status_var.set("Fix burn-in inputs and try again.")
            return

        test_hw = None
        try:
            from burnin_runtime import BurnInHardware
            test_hw = BurnInHardware()
            test_hw.read_voltage(self.burnin_config.pad1)
            test_hw.read_voltage(self.burnin_config.pad2)
            
        except Exception as exc:
            self.burnin_status_var.set("Burn-in hardware check failed.")
            messagebox.showerror(
                "Burn-In hardware Failed",
                f"Could not communicate with the burn-in PCB.\n\n"
                f"Make sure the burn-in PCB is connected.\n\n"
                f"Error:\n{exc}"
            )

        self.burnin_stop_event.clear()
        self.burnin_raw_rows.clear()
        self.burnin_rows.clear()
        self.burnin_cycle_buffers.clear()
        self.burnin_completed_cycles.clear()
        self.burnin_x.clear()
        self.burnin_y.clear()
        self.burnin_final_status = "RUNNING"
        self.burnin_failure_reason = ""
        self.reset_burnin_plot()

        serial = self.serial_var.get().strip().upper()
        phase = self.phase_var.get().strip()
        tester = self.tester_var.get().strip()
        cfg = self.burnin_config

        self.burnin_running = True
        self.burnin_start_btn.configure(state="disabled")
        self.burnin_stop_btn.configure(state="normal")
        self.burnin_status_var.set("Burn-in running...")
        self.burnin_log("Start Burn-In pressed.")
        self.burnin_log(
            f"Serial: {serial} | Phase: {phase} | Tester: {tester} | Pads: {cfg.pad1}-{cfg.pad2} | "
            f"Cycles: {cfg.pulses} | On/Off: {cfg.on_time}/{cfg.off_time}s | Threshold: ±{PERCENT_THRESHOLD:.1f}% | "
            f"GUI averaging: first {GUI_AVERAGE_SAMPLES_PER_CYCLE} samples/cycle"
        )

        self.burnin_thread = threading.Thread(target=self.burnin_worker, daemon=True)
        self.burnin_thread.start()

    def burnin_worker(self) -> None:
        try:
            for row in run_burnin(self.burnin_config, stop_requested=self.burnin_stop_event.is_set):
                self.burnin_queue.put(("sample", row))
            self.burnin_queue.put(("done", "COMPLETE", ""))
        except BurnInError as exc:
            self.burnin_queue.put(("done", "FAILED", str(exc)))
        except Exception as exc:
            details = f"{exc}\n\n{traceback.format_exc()}"
            self.burnin_queue.put(("done", "ERROR", details))

    def process_burnin_queue(self) -> None:
        try:
            while True:
                msg = self.burnin_queue.get_nowait()
                msg_type = msg[0]

                if msg_type == "sample":
                    self.handle_burnin_sample(msg[1])

                elif msg_type == "done":
                    self.finalize_all_pending_cycle_averages()
                    self.finish_burnin(msg[1], msg[2])

        except queue.Empty:
            pass
        except Exception as exc:
            # Never let a GUI-side update error kill Tkinter silently.
            self.burnin_running = False
            self.burnin_start_btn.configure(state="normal")
            self.burnin_stop_btn.configure(state="disabled")
            self.burnin_status_var.set("Burn-in GUI update error.")
            self.burnin_log(f"GUI queue/update error: {exc}")
            messagebox.showerror("Burn-In GUI Error", str(exc))

        self.after(50, self.process_burnin_queue)

    def handle_burnin_sample(self, row: dict) -> None:
        """
        Receive raw streamed measurements from burnin_runtime.

        burnin_runtime is intentionally left unchanged and may yield many raw
        measurements per cycle. The GUI reduces those raw rows to one averaged
        point per cycle for plotting/export.
        """
        self.burnin_raw_rows.append(row)

        cycle = int(row.get("cycle", 0))
        resistance = float(row.get("resistance_ohms", 0.0))
        status = row.get("status", "")

        if status == "BASELINE" or cycle == 0:
            if not any(r.get("status") == "BASELINE" for r in self.burnin_rows):
                self.burnin_rows.append(row)
            self.burnin_log(f"Baseline resistance: {resistance:.6f} Ω")
            return

        self.burnin_status_var.set(f"Running... Cycle {cycle} of {self.burnin_config.pulses}")

        if cycle in self.burnin_completed_cycles:
            return

        buffer = self.burnin_cycle_buffers.setdefault(cycle, [])
        buffer.append(row)

        # Plot/export only one point per cycle: average of first N raw samples.
        if len(buffer) >= GUI_AVERAGE_SAMPLES_PER_CYCLE:
            self.finalize_cycle_average(cycle)

    def finalize_cycle_average(self, cycle: int) -> None:
        if cycle in self.burnin_completed_cycles:
            return

        samples = self.burnin_cycle_buffers.get(cycle, [])
        if not samples:
            return

        selected = samples[:GUI_AVERAGE_SAMPLES_PER_CYCLE]
        n = len(selected)

        avg_resistance = sum(float(r.get("resistance_ohms", 0.0)) for r in selected) / n
        avg_percent = sum(float(r.get("percent_change", 0.0)) for r in selected) / n
        last = selected[-1]

        avg_row = dict(last)
        avg_row["sample_index"] = cycle
        avg_row["resistance_ohms"] = avg_resistance
        avg_row["percent_change"] = avg_percent
        avg_row["status"] = f"AVG_FIRST_{n}_SAMPLES"

        self.burnin_rows.append(avg_row)
        self.burnin_completed_cycles.add(cycle)

        self.burnin_x.append(cycle)
        self.burnin_y.append(avg_resistance)
        self.update_burnin_plot()

        self.burnin_log(
            f"Cycle {cycle}: average of first {n} samples = "
            f"{avg_resistance:.6f} Ω ({avg_percent:.2f}% from baseline)"
        )

    def finalize_all_pending_cycle_averages(self) -> None:
        for cycle in sorted(self.burnin_cycle_buffers):
            self.finalize_cycle_average(cycle)

    def finish_burnin(self, status: str, reason: str) -> None:
        self.burnin_running = False
        self.burnin_start_btn.configure(state="normal")
        self.burnin_stop_btn.configure(state="disabled")
        self.burnin_final_status = status
        self.burnin_failure_reason = reason

        if status == "COMPLETE":
            self.burnin_status_var.set("Burn-in complete.")
            self.burnin_log("Burn-in complete.")
        elif status == "FAILED":
            self.burnin_status_var.set("Burn-in failed safely. Current disabled.")
            self.burnin_log(f"FAILED: {reason}")
            messagebox.showerror("Burn-In Failed", reason)
        else:
            self.burnin_status_var.set("Burn-in error. Current disabled.")
            self.burnin_log(f"ERROR: {reason}")
            messagebox.showerror("Burn-In Error", reason)

        self.export_current_burnin_run()

    def on_burnin_stop(self) -> None:
        if self.burnin_running:
            self.burnin_stop_event.set()
            self.burnin_status_var.set("Stopping burn-in safely...")
            self.burnin_log("Stop requested. Current will be disabled by burn-in runtime.")
        else:
            self.burnin_status_var.set("Stopped.")
            self.burnin_log("Stop pressed while burn-in was not running.")

    def reset_burnin_plot(self) -> None:
        if not MATPLOTLIB_AVAILABLE:
            return
        self.burnin_line.set_data([], [])
        cfg = self.burnin_config
        serial = self.serial_var.get().strip().upper()
        phase = self.phase_var.get().strip()
        tester = self.tester_var.get().strip()
        if cfg:
            title = f"{serial} | {phase} | {tester} | Pads {cfg.pad1}-{cfg.pad2} | {cfg.current_amps} A"
        else:
            title = "Burn-In Resistance"
        self.burnin_ax.set_title(title)
        self.burnin_ax.set_xlabel("Cycle #")
        self.burnin_ax.set_ylabel("Resistance (Ω)")
        self.burnin_ax.relim()
        self.burnin_ax.autoscale_view()
        self.burnin_canvas.draw_idle()

    def update_burnin_plot(self) -> None:
        if not MATPLOTLIB_AVAILABLE:
            return
        self.burnin_line.set_data(self.burnin_x, self.burnin_y)
        self.burnin_ax.relim()
        self.burnin_ax.autoscale_view()
        self.burnin_canvas.draw_idle()

    def export_current_burnin_run(self) -> None:
        if not self.burnin_config or not self.burnin_rows:
            return

        try:
            wb_path = export_burnin_to_excel(
                outdir=self.output_dir_var.get().strip(),
                serial=self.serial_var.get().strip().upper(),
                phase=self.phase_var.get().strip(),
                tester_name=self.tester_var.get().strip(),
                notes=self.burnin_metadata.get_notes(),
                burnin_config=self.burnin_config,
                rows=self.burnin_rows,
                final_status=self.burnin_final_status,
                failure_reason=self.burnin_failure_reason,
            )
            self.burnin_log(f"Exported burn-in run to: {wb_path}")
            open_file(wb_path)
            self.refresh_serial_list()
        except Exception as exc:
            self.burnin_log(f"ERROR exporting burn-in run: {exc}")
            messagebox.showerror("Burn-In Export Failed", str(exc))


def main() -> None:
    root = tk.Tk()
    root.title("QForge Test Controller")
    root.minsize(980, 800)
    root.columnconfigure(0, weight=1)
    root.rowconfigure(0, weight=1)
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
